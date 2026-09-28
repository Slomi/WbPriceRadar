"""Тексты карточек и отчётов, график истории цен, выгрузка в Excel."""
import io
from datetime import datetime, timedelta, timezone
from html import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import db  # noqa: E402
from utils import local, now, rub, short  # noqa: E402

ROLE = {"mine": "⭐ Ваш товар", "rival": "🎯 Конкурент"}


def _url(nm: int) -> str:
    return f"https://www.wildberries.ru/catalog/{nm}/detail.aspx"


async def change(nm: int, price: float | None, days: int) -> str:
    if not price:
        return "—"
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    old = await db.price_between(nm, (cutoff - timedelta(days=days)).isoformat(), cutoff.isoformat())
    if not old:
        return "—"
    pct = (price - old) / old * 100
    return "без изменений" if abs(pct) < 0.5 else f"{pct:+.0f}%"


def _diff(price: float | None, base: float | None) -> str:
    if not price or not base:
        return ""
    d = price - base
    if abs(d) < 1:
        return " · = вашей"
    return f" · <b>на {rub(-d)} дешевле</b> ❗" if d < 0 else f" · на {rub(d)} дороже"


async def card_text(p) -> str:
    lines = [
        ROLE[p["role"]],
        f"<b>{escape(p['name'] or '')}</b>",
        f"{escape(p['brand'] or '—')} · продавец {escape(p['supplier'] or '—')}",
        f'Артикул <a href="{_url(p["nm"])}">{p["nm"]}</a>',
        "",
    ]
    if p["price"]:
        disc = round((1 - p["price"] / p["basic"]) * 100) if p["basic"] else 0
        lines.append(f"💰 <b>{rub(p['price'])}</b>" + (f"  (до скидки {rub(p['basic'])}, −{disc}%)" if disc else ""))
    else:
        lines.append("💰 нет в наличии")
    lines += [
        f"📦 Остаток: {p['qty']} шт",
        f"⭐ {p['rating']:.1f} · отзывов: {p['feedbacks']:,}".replace(",", " "),
        f"📊 За 24 ч: {await change(p['nm'], p['price'], 1)} · за 7 дней: {await change(p['nm'], p['price'], 7)}",
    ]
    if p["role"] == "mine":
        rivals = await db.rivals_of(p["id"])
        if rivals:
            cheaper = [r for r in rivals if r["price"] and p["price"] and r["price"] < p["price"]]
            lines += ["", f"<b>Конкуренты ({len(rivals)})</b>" + (f", дешевле вас: {len(cheaper)}" if cheaper else "")]
            lines += [f"• {escape(short(r['name'], 32))} — {rub(r['price'])}{_diff(r['price'], p['price'])}" for r in rivals]
    elif p["parent_id"]:
        mine = await db.get_product(p["parent_id"], p["user_id"])
        if mine:
            lines += ["", f"Ваш товар «{escape(short(mine['name'], 30))}»: {rub(mine['price'])}"
                      + _diff(p["price"], mine["price"])]
    if p["checked_at"]:
        lines += ["", f"<i>Проверено {local(p['checked_at']):%d.%m %H:%M}</i>"]
    return "\n".join(lines)


async def daily_report(user_id: int) -> str:
    products = await db.user_products(user_id)
    if not products:
        return "Пока нечего отслеживать — добавьте товары кнопкой «➕ Добавить товар»."
    mine = [p for p in products if p["role"] == "mine"]
    loose = [p for p in products if p["role"] == "rival" and not p["parent_id"]]
    out = [f"📊 <b>Сводка на {now():%d.%m %H:%M}</b>"]
    alerts = 0
    for m in mine:
        rivals = await db.rivals_of(m["id"])
        block = [f"\n⭐ <b>{escape(short(m['name'], 45))}</b>",
                 f"Ваша цена: <b>{rub(m['price'])}</b> (за сутки: {await change(m['nm'], m['price'], 1)}), "
                 f"остаток {m['qty']} шт"]
        for r in rivals:
            mark = ""
            if r["price"] and m["price"] and r["price"] < m["price"]:
                mark, alerts = " ❗", alerts + 1
            if r["qty"] == 0:
                mark += " ❌ нет в наличии"
            block.append(f"  🎯 {escape(short(r['name'], 32))}: {rub(r['price'])}"
                         f" ({await change(r['nm'], r['price'], 1)}){mark}")
        if rivals:
            priced = [r for r in rivals if r["price"]]
            if priced and m["price"]:
                cheapest = min(priced, key=lambda r: r["price"])
                pos = 1 + sum(1 for r in priced if r["price"] < m["price"])
                block.append(f"  Ваше место по цене: {pos} из {len(priced) + 1}"
                             + (f", минимальная у конкурентов — {rub(cheapest['price'])}" if cheapest["price"] < m["price"] else " — вы самые дешёвые 👍"))
        out += block
    if loose:
        out.append("\n🎯 <b>Конкуренты без привязки</b>")
        out += [f"  {escape(short(r['name'], 40))}: {rub(r['price'])} ({await change(r['nm'], r['price'], 1)})" for r in loose]
    if alerts:
        out.insert(1, f"❗ Конкурентов дешевле вас: <b>{alerts}</b>")
    return "\n".join(out)


def _series(rows) -> tuple[list[datetime], list[float]]:
    pts = sorted((local(r["ts"]), r["price"]) for r in rows if r["price"])
    return [p[0] for p in pts], [p[1] for p in pts]


async def chart_png(p, days: int = 90) -> bytes:
    since = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()
    if p["role"] == "mine":
        main, others = p, await db.rivals_of(p["id"])
    else:
        mine = await db.get_product(p["parent_id"], p["user_id"]) if p["parent_id"] else None
        main, others = (mine, [p]) if mine else (p, [])

    fig, ax = plt.subplots(figsize=(10, 5.4), dpi=130)
    fig.patch.set_facecolor("#ffffff")
    palette = ["#e8833a", "#8e6cd8", "#2fa58c", "#d24d57", "#6b7a8f"]
    for i, r in enumerate(others[:5]):
        xs, ys = _series(await db.history(r["nm"], since))
        if xs:
            ax.step(xs, ys, where="post", color=palette[i], lw=1.8, alpha=0.9, label=f"{short(r['name'], 24)} · {rub(r['price'])}")
    xs, ys = _series(await db.history(main["nm"], since))
    if xs:
        label = ("Вы: " if main["role"] == "mine" else "") + f"{short(main['name'], 24)} · {rub(main['price'])}"
        ax.step(xs, ys, where="post", color="#2a7de1", lw=3, label=label)
        ax.scatter([xs[-1]], [ys[-1]], color="#2a7de1", zorder=5)
        ax.annotate(rub(ys[-1]), (xs[-1], ys[-1]), textcoords="offset points", xytext=(6, 6), fontsize=10,
                    color="#2a7de1", fontweight="bold")
    ax.set_title(f"Цена за {days} дней" + (" — вы и конкуренты" if others else ""), fontsize=14, loc="left")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:,.0f} ₽".replace(",", " ")))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d.%m"))
    ax.grid(alpha=0.25)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    handles, labels = ax.get_legend_handles_labels()  # «вы» — первым в легенде
    order = sorted(range(len(labels)), key=lambda i: not labels[i].startswith("Вы: "))
    ax.legend([handles[i] for i in order], [labels[i] for i in order], title="Конкуренты и вы" if others else None,
              loc="upper left", bbox_to_anchor=(1.01, 1), frameon=False, fontsize=9)
    ax.text(0, -0.13, "До сегодняшнего дня — недельная история цен WB, дальше — проверки бота",
            transform=ax.transAxes, fontsize=8, color="#888")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


async def excel(user_id: int) -> bytes:
    products = await db.user_products(user_id)
    by_id = {p["id"]: p for p in products}
    wbk = Workbook()
    ws = wbk.active
    ws.title = "Товары"
    head = ["Роль", "Мой товар", "Артикул", "Название", "Бренд", "Продавец", "Цена, ₽", "До скидки, ₽", "Скидка, %",
            "Разница с моим, ₽", "Остаток, шт", "Рейтинг", "Отзывы", "За 24 ч", "За 7 дней", "Проверено", "Ссылка"]
    ws.append(head)
    for p in products:
        parent = by_id.get(p["parent_id"]) if p["parent_id"] else None
        disc = round((1 - p["price"] / p["basic"]) * 100) if p["price"] and p["basic"] else None
        diff = round(p["price"] - parent["price"]) if parent and p["price"] and parent["price"] else None
        ws.append([
            "Мой" if p["role"] == "mine" else "Конкурент", parent["name"] if parent else "", p["nm"], p["name"],
            p["brand"], p["supplier"], p["price"], p["basic"], disc, diff, p["qty"], p["rating"], p["feedbacks"],
            await change(p["nm"], p["price"], 1), await change(p["nm"], p["price"], 7),
            local(p["checked_at"]).strftime("%d.%m.%Y %H:%M") if p["checked_at"] else "", _url(p["nm"]),
        ])
    hs = wbk.create_sheet("История")
    hs.append(["Артикул", "Название", "Дата", "Цена, ₽", "Остаток, шт", "Источник"])
    for p in products:
        for h in await db.history(p["nm"]):
            when = local(h["ts"]).strftime("%d.%m.%Y" if h["source"] == "wb" else "%d.%m.%Y %H:%M")
            hs.append([p["nm"], p["name"], when, h["price"], h["qty"],
                       "WB (неделя)" if h["source"] == "wb" else "проверка бота"])
    for sheet in (ws, hs):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2A7DE1")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col[:200])
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 45)
    for row in ws.iter_rows(min_row=2):
        if row[9].value is not None and row[9].value < 0:
            for c in row:
                c.fill = PatternFill("solid", fgColor="FDE2E2")  # конкурент дешевле моего
    buf = io.BytesIO()
    wbk.save(buf)
    return buf.getvalue()
