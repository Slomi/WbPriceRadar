"""Прогон демо-сценария через НАСТОЯЩИЕ обработчики aiogram 3 без Telegram.

Подменяется только транспорт: всё, что бот отправил бы в Telegram (сообщения, правки, кнопки, фото,
документы, всплывашки), записывается в script.json, из которого Remotion-шаблон собирает ролик.

Использование (из video/gen_script.py проекта):

    demo = TgDemo(routers=[admin.router, client.router], bot_name="Бритва",
                  chats={2000: ("Иван", "Клиент"), 1000: ("Мастер", "Мастер / админ")},
                  media_dir=VIDEO / "public" / "media")
    await demo.say(2000, "/start")
    await demo.click(2000, "✂️")                    # последняя inline-кнопка, чей текст начинается так
    await demo.say(1000, "Ответ", reply_to=demo.last_bot(1000, "🆘"))
    demo.clock("за 2 часа до визита")               # заставка «прошло время» в ролике
    await some_background_job(demo.bot)             # напоминания/мониторинг — настоящими функциями бота
    demo.save(VIDEO / "src" / "data" / "script.json")

Проект сам готовит окружение ДО импорта своих модулей: временная БД, ADMIN_IDS = id из chats и т.п.
"""
from __future__ import annotations

import json
import time
from pathlib import Path

from aiogram import Bot, Dispatcher, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.enums import ParseMode
from aiogram.types import (InlineKeyboardMarkup, Message, MessageId, ReplyKeyboardMarkup,
                           ReplyKeyboardRemove, Update)

BOT_ID = 9999


def _markup(m) -> dict | None:
    if isinstance(m, InlineKeyboardMarkup):
        return {"inline": [[b.text for b in row] for row in m.inline_keyboard]}
    if isinstance(m, ReplyKeyboardMarkup):
        return {"reply": [[b.text for b in row] for row in m.keyboard]}
    if isinstance(m, ReplyKeyboardRemove):
        return {"remove": True}
    return None


def _xlsx_preview(path: Path, rows: int = 7, cols: int = 7) -> dict | None:
    """Первые строки первого листа — чтобы показать «что внутри Excel» в ролике."""
    try:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).worksheets[0]
        data = [["" if v is None else str(v) for v in row[:cols]]
                for row in ws.iter_rows(max_row=rows, values_only=True)]
        return {"sheet": ws.title, "rows": data}
    except Exception:
        return None


class _Session(BaseSession):
    def __init__(self, demo: "TgDemo"):
        super().__init__()
        self.demo = demo

    async def make_request(self, bot, method, timeout=None):
        d, name = self.demo, type(method).__name__
        if name == "SendMessage":
            mid = d.new_id()
            inline = method.reply_markup if isinstance(method.reply_markup, InlineKeyboardMarkup) else None
            m, raw = d.message(method.chat_id, mid, method.text, markup=inline)
            d.raw[(method.chat_id, mid)] = raw
            d.events.append({"op": "bot", "chat": method.chat_id, "id": mid, "text": method.text,
                             "markup": _markup(method.reply_markup)})
            return m
        if name == "EditMessageText":
            raw = d.raw[(method.chat_id, method.message_id)]
            raw["text"] = method.text
            if isinstance(method.reply_markup, InlineKeyboardMarkup):
                raw["reply_markup"] = method.reply_markup.model_dump(exclude_none=True)
            else:
                raw.pop("reply_markup", None)
            d.events.append({"op": "edit", "chat": method.chat_id, "id": method.message_id, "text": method.text,
                             "markup": _markup(method.reply_markup)})
            return d.message(method.chat_id, method.message_id, method.text)[0]
        if name == "EditMessageReplyMarkup":
            raw = d.raw[(method.chat_id, method.message_id)]
            if isinstance(method.reply_markup, InlineKeyboardMarkup):
                raw["reply_markup"] = method.reply_markup.model_dump(exclude_none=True)
            else:
                raw.pop("reply_markup", None)
            d.events.append({"op": "edit_markup", "chat": method.chat_id, "id": method.message_id,
                             "markup": _markup(method.reply_markup)})
            return d.message(method.chat_id, method.message_id, raw.get("text", ""))[0]
        if name == "CopyMessage":
            src = d.raw[(method.from_chat_id, method.message_id)]
            mid = d.new_id()
            d.events.append({"op": "bot", "chat": method.chat_id, "id": mid, "text": src.get("text", ""),
                             "markup": None})
            return MessageId(message_id=mid)
        if name in ("SendPhoto", "SendDocument"):
            mid = d.new_id()
            media = method.photo if name == "SendPhoto" else method.document
            fname = f"{mid}_{getattr(media, 'filename', None) or 'file'}"
            (d.media_dir / fname).write_bytes(media.data)
            ev = {"op": "photo" if name == "SendPhoto" else "document", "chat": method.chat_id, "id": mid,
                  "file": fname, "text": method.caption or "", "markup": _markup(method.reply_markup)}
            if fname.lower().endswith(".xlsx"):
                ev["preview"] = _xlsx_preview(d.media_dir / fname)
            d.events.append(ev)
            m, raw = d.message(method.chat_id, mid, method.caption or "")
            d.raw[(method.chat_id, mid)] = raw
            return m
        if name == "AnswerCallbackQuery":
            if method.text:
                d.events.append({"op": "toast", "chat": None, "text": method.text})
            return True
        if name in ("SendChatAction", "SetMyCommands", "DeleteWebhook", "DeleteMessage"):
            return True
        raise RuntimeError(f"tg_mock: метод {name} не поддержан — добавьте его в _Session.make_request")

    async def close(self):
        pass

    async def stream_content(self, *a, **kw):
        raise RuntimeError("tg_mock: скачивание файлов не поддержано")


class TgDemo:
    def __init__(self, routers: list[Router], bot_name: str, chats: dict[int, tuple[str, str]], media_dir: Path):
        """chats: {user_id: (имя пользователя в Telegram, подпись окна в ролике)} — 1 или 2 окна."""
        self.events: list[dict] = []
        self.raw: dict[tuple[int, int], dict] = {}
        self.chats = chats
        self.bot_user = {"id": BOT_ID, "is_bot": True, "first_name": bot_name}
        self.users = {uid: {"id": uid, "is_bot": False, "first_name": name} for uid, (name, _) in chats.items()}
        self.media_dir = Path(media_dir)
        self.media_dir.mkdir(parents=True, exist_ok=True)
        self._id, self._upd = 500, 0
        self.bot = Bot("1:demo", session=_Session(self), default=DefaultBotProperties(parse_mode=ParseMode.HTML))
        self.dp = Dispatcher()
        self.dp.include_routers(*routers)

    def new_id(self) -> int:
        self._id += 1
        return self._id

    def message(self, chat_id: int, mid: int, text: str, from_user: dict | None = None, markup=None,
                reply_to: dict | None = None):
        raw = {"message_id": mid, "date": int(time.time()), "chat": {"id": chat_id, "type": "private"},
               "from": from_user or self.bot_user, "text": text}
        if markup is not None:
            raw["reply_markup"] = markup.model_dump(exclude_none=True)
        if reply_to:
            raw["reply_to_message"] = reply_to
        return Message.model_validate(raw, context={"bot": self.bot}), raw

    async def _feed(self, payload: dict):
        self._upd += 1
        await self.dp.feed_update(self.bot, Update.model_validate({"update_id": self._upd, **payload},
                                                                  context={"bot": self.bot}))

    async def say(self, uid: int, text: str, reply_to: int | None = None):
        """Пользователь пишет текст (или жмёт кнопку reply-клавиатуры — это тоже текст)."""
        mid = self.new_id()
        _, raw = self.message(uid, mid, text, from_user=self.users[uid],
                              reply_to=self.raw.get((uid, reply_to)) if reply_to else None)
        self.raw[(uid, mid)] = raw
        self.events.append({"op": "user", "chat": uid, "id": mid, "text": text, "reply_to": reply_to})
        await self._feed({"message": raw})

    async def click(self, uid: int, prefix: str) -> bool:
        """Жмёт inline-кнопку в самом свежем сообщении, где есть кнопка с таким началом текста."""
        for (chat, mid), raw in sorted(self.raw.items(), key=lambda kv: -kv[0][1]):
            if chat != uid:
                continue
            for row in raw.get("reply_markup", {}).get("inline_keyboard", []):
                for b in row:
                    if b["text"].startswith(prefix):
                        self.events.append({"op": "click", "chat": uid, "id": mid, "button": b["text"]})
                        await self._feed({"callback_query": {"id": str(self._upd), "from": self.users[uid],
                                                             "chat_instance": "demo", "message": raw,
                                                             "data": b["callback_data"]}})
                        return True
        return False

    def last_bot(self, uid: int, starts: str = "") -> int:
        """id последнего сообщения бота в чате uid, текст которого начинается с starts (для reply_to)."""
        return next(e["id"] for e in reversed(self.events)
                    if e.get("chat") == uid and e["op"] == "bot" and e.get("text", "").startswith(starts))

    def clock(self, text: str):
        """Заставка «прошло время» — перед напоминаниями, плановыми проверками и т.п."""
        self.events.append({"op": "clock", "chat": None, "text": text})

    def save(self, path: Path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        data = {"chats": [{"id": uid, "label": label} for uid, (_, label) in self.chats.items()],
                "events": self.events}
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"событий: {len(self.events)} → {path}")
